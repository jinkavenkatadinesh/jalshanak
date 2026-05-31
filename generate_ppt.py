import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Dark Ocean Theme)
BG_COLOR = RGBColor(15, 23, 42)       # Deep slate navy (#0F172A)
CARD_BG_COLOR = RGBColor(30, 41, 59)  # Lightened slate for cards (#1E293B)
CARD_BORDER = RGBColor(51, 65, 85)    # Slate border (#334155)
TEXT_WHITE = RGBColor(248, 250, 252)  # Ice white (#F8FAFC)
TEXT_MUTED = RGBColor(148, 163, 184)  # Slate gray (#94A3B8)
ACCENT_TEAL = RGBColor(6, 182, 212)   # Vibrant teal (#06B6D4)
ACCENT_SKY = RGBColor(56, 189, 248)   # Bright sky blue (#38BDF8)

blank_layout = prs.slide_layouts[6]

def apply_slide_base(slide, category_name, title_text):
    """
    Applies deep dark background, top header, and horizontal rules to a content slide.
    """
    # 1. Base Dark Background
    bg = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE = 1
        0, 0, prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()  # No border
    
    # 2. Category Pill/Label
    cat_box = slide.shapes.add_textbox(Inches(0.866), Inches(0.4), Inches(8), Inches(0.4))
    tf_cat = cat_box.text_frame
    tf_cat.word_wrap = True
    tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_name.upper()
    p_cat.font.name = 'Trebuchet MS'
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = ACCENT_TEAL
    
    # 3. Main Header Title
    title_box = slide.shapes.add_textbox(Inches(0.866), Inches(0.7), Inches(10), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = 'Trebuchet MS'
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE
    
    # 4. Underline accent
    line = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(0.866), Inches(1.5), Inches(2.0), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_TEAL
    line.line.fill.background()

def create_card(slide, left, top, width, height, title, body_text):
    """
    Creates an elegant dark slate card with custom border and highly formatted text block.
    """
    # Card Background container
    card = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG_COLOR
    card.line.color.rgb = CARD_BORDER
    card.line.width = Pt(1.5)
    
    # Content Textbox
    textbox = slide.shapes.add_textbox(
        left + Inches(0.2), 
        top + Inches(0.2), 
        width - Inches(0.4), 
        height - Inches(0.4)
    )
    tf = textbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    # Card Title
    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.name = 'Trebuchet MS'
    p_title.font.size = Pt(18)
    p_title.font.bold = True
    p_title.font.color.rgb = ACCENT_SKY
    p_title.space_after = Pt(12)
    
    # Card Body
    p_body = tf.add_paragraph()
    p_body.text = body_text
    p_body.font.name = 'Calibri'
    p_body.font.size = Pt(13)
    p_body.font.color.rgb = TEXT_MUTED
    p_body.line_spacing = 1.2
    
    return card

# ----------------- SLIDE 1: TITLE SLIDE -----------------
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
bg1.fill.solid()
bg1.fill.fore_color.rgb = BG_COLOR
bg1.line.fill.background()

# Title text box
title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.0))
tf1 = title_box.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

p1 = tf1.paragraphs[0]
p1.text = "JalRakshak"
p1.font.name = 'Trebuchet MS'
p1.font.size = Pt(56)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE
p1.space_after = Pt(6)

p2 = tf1.add_paragraph()
p2.text = "Smart Water Leak Reporting System"
p2.font.name = 'Trebuchet MS'
p2.font.size = Pt(28)
p2.font.bold = True
p2.font.color.rgb = ACCENT_TEAL
p2.space_after = Pt(24)

# Decorative Divider
divider = slide1.shapes.add_shape(1, Inches(1.0), Inches(4.3), Inches(3.0), Inches(0.06))
divider.fill.solid()
divider.fill.fore_color.rgb = ACCENT_SKY
divider.line.fill.background()

# Presenter/Meta Box
meta_box = slide1.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(8), Inches(1.0))
tf_meta = meta_box.text_frame
p_meta1 = tf_meta.paragraphs[0]
p_meta1.text = "Empowering Hyderabad Citizens & Streamlining Civic Water Infrastructure"
p_meta1.font.name = 'Calibri'
p_meta1.font.size = Pt(14)
p_meta1.font.color.rgb = TEXT_MUTED
p_meta1.font.italic = True

p_meta2 = tf_meta.add_paragraph()
p_meta2.text = "CivicTech Solutions | Hyderabad & Telangana Metropolitan Division"
p_meta2.font.name = 'Calibri'
p_meta2.font.size = Pt(13)
p_meta2.font.color.rgb = TEXT_MUTED
p_meta2.space_before = Pt(4)


# ----------------- SLIDE 2: THE CIVIC CHALLENGE -----------------
slide2 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide2, "Civic Challenge", "The Water Leakage Challenge")

# 3 Column layout
left_margin = Inches(0.866)
card_width = Inches(3.6)
gap = Inches(0.4)
top_pos = Inches(2.0)
card_height = Inches(4.5)

create_card(
    slide2, left_margin, top_pos, card_width, card_height,
    "Unreported Waste",
    "Thousands of gallons of clean water are lost daily in municipal divisions due to slow reporting and lag in civic repair cycles."
)
create_card(
    slide2, left_margin + card_width + gap, top_pos, card_width, card_height,
    "Citizen Friction",
    "Municipal divisions lack a direct, real-time channels for citizens to notify authorities instantly with precise geographic telemetry."
)
create_card(
    slide2, left_margin + (card_width + gap) * 2, top_pos, card_width, card_height,
    "Data Gap",
    "Civic boards struggle to coordinate resolutions, analyze hotspots, or weed out duplicate tickets because of offline workflows."
)


# ----------------- SLIDE 3: THE JALRAKSHAK SOLUTION -----------------
slide3 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide3, "Core Solution", "Empowering Civic Action")

create_card(
    slide3, left_margin, top_pos, card_width, card_height,
    "The Platform",
    "A state-of-the-art CivicTech full-stack web application designed to securely connect active citizens directly with municipal water authorities."
)
create_card(
    slide3, left_margin + card_width + gap, top_pos, card_width, card_height,
    "Reporting Portal",
    "A fluid interface offering 1-click citizen reporting with automated location GPS snapping and instant photo proof uploads."
)
create_card(
    slide3, left_margin + (card_width + gap) * 2, top_pos, card_width, card_height,
    "Command Center",
    "An advanced analytical portal with real-time KPI aggregates, status transition workflows, and automated issue severity tracking."
)


# ----------------- SLIDE 4: HIGH-PERFORMANCE TECH STACK -----------------
slide4 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide4, "System Architecture", "Unified Technology Stack")

create_card(
    slide4, left_margin, top_pos, card_width, card_height,
    "Frontend Layer",
    "Built with React.js (Vite) for performance, Leaflet.js maps with custom OpenStreetMap canvases, and styled with premium Vanilla CSS Glassmorphism."
)
create_card(
    slide4, left_margin + card_width + gap, top_pos, card_width, card_height,
    "Backend Core",
    "Powered by high-performance FastAPI, SQLAlchemy ORM, robust Pydantic v2 data validation schemas, and a custom Python Heuristics Engine."
)
create_card(
    slide4, left_margin + (card_width + gap) * 2, top_pos, card_width, card_height,
    "DevOps & Data",
    "Structured for PostgreSQL (production) with an SQLite fallback sandbox. Orchestrated via Docker containers behind a secure Nginx reverse proxy gateway."
)


# ----------------- SLIDE 5: CITIZEN REPORTING EXPERIENCE -----------------
slide5 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide5, "Citizen Experience", "Streamlined Citizen Features")

# 4 Columns layout
card_w4 = Inches(2.6)
gap_w4 = Inches(0.4)
create_card(
    slide5, left_margin, top_pos, card_w4, card_height,
    "Auto-GPS",
    "Instantly captures precise telemetry coordinates via browser Geolocation API and centers Hyderabad custom canvas pins."
)
create_card(
    slide5, left_margin + card_w4 + gap_w4, top_pos, card_w4, card_height,
    "Photo Proof",
    "Seamless drag-and-drop or click loaders for attaching visual evidence of public water leakage reports."
)
create_card(
    slide5, left_margin + (card_w4 + gap_w4) * 2, top_pos, card_w4, card_height,
    "Duplicate Check",
    "Checks a 100-meter radius dynamically in the backend and warns citizens if another active report already exists nearby."
)
create_card(
    slide5, left_margin + (card_w4 + gap_w4) * 3, top_pos, card_w4, card_height,
    "P2P Voting",
    "Citizens can search the map and vote to verify reported leaks, boosting report credibility metrics for speedy resolutions."
)


# ----------------- SLIDE 6: AUTHORITY ANALYTICS DASHBOARD -----------------
slide6 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide6, "Authority Portal", "Data-Driven Command Center")

create_card(
    slide6, left_margin, top_pos, card_w4, card_height,
    "Real-Time KPIs",
    "Visual indicators counting active, pending, and resolved leakages to keep teams informed at a glance."
)
create_card(
    slide6, left_margin + card_w4 + gap_w4, top_pos, card_w4, card_height,
    "Visual Charts",
    "Dynamic distribution charts displaying area-wise issue densities, current status splits, and ticket severity breakdowns."
)
create_card(
    slide6, left_margin + (card_w4 + gap_w4) * 2, top_pos, card_w4, card_height,
    "Data Grid",
    "Multi-column sorting tables linked to a split-view inspector and sliding side-drawer for in-depth ticket audits."
)
create_card(
    slide6, left_margin + (card_w4 + gap_w4) * 3, top_pos, card_w4, card_height,
    "Timelines",
    "Interactive status logs that record detailed technician diagnostic remarks across transitions (Reported ➔ In Progress ➔ Resolved)."
)


# ----------------- SLIDE 7: UNDER THE HOOD: SMART ENGINE -----------------
slide7 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide7, "System Engineering", "Intelligent Civic Engineering")

create_card(
    slide7, left_margin, top_pos, card_width, card_height,
    "Haversine merges",
    "Uses a proximity-checking Haversine distance formula to identify duplicates and dynamically merge reports to prevent database clutter."
)
create_card(
    slide7, left_margin + card_width + gap, top_pos, card_width, card_height,
    "Priority Scorer",
    "Calculates severity metrics dynamically based on initial hazard rank, ticket duration, and cumulative citizen confirmation votes."
)
create_card(
    slide7, left_margin + (card_width + gap) * 2, top_pos, card_width, card_height,
    "AI Diagnostics",
    "Performs automatic description scanning to flag leak severity and logs details in administrative inspection files."
)


# ----------------- SLIDE 8: FUTURE ROADMAP & IMPACT -----------------
slide8 = prs.slides.add_slide(blank_layout)
apply_slide_base(slide8, "Future Vision", "Future Vision & Project Impact")

create_card(
    slide8, left_margin, top_pos, card_width, card_height,
    "Immediate Impact",
    "Minimizes civic water loss, shortens administrative repair cycles, empowers active citizens, and builds high civic trust."
)
create_card(
    slide8, left_margin + card_width + gap, top_pos, card_width, card_height,
    "Vision 2026",
    "Incorporate computer vision AI models to automatically estimate leak flow rates and size of damage from uploaded photo evidence."
)
create_card(
    slide8, left_margin + (card_width + gap) * 2, top_pos, card_width, card_height,
    "Scale Features",
    "Deploy multi-channel notification systems (SMS, WhatsApp) for citizen status updates, alongside leaderboard-driven reporter gamification."
)

# Save presentation
prs.save("JalRakshak_Project_Presentation.pptx")
print("Presentation compiled successfully as 'JalRakshak_Project_Presentation.pptx'")
